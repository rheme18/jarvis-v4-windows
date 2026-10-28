# -*- coding: utf-8 -*-
"""
J.A.R.V.I.S V4 — Office & Project Automation Suite
Voice & Tool Driven Document Creation:
1. Excel Tabloları (openpyxl) — Profesyonel renk paletleri, kenarlıklar, formüller ve otomatik sütun genişliği.
2. PowerPoint Sunumları (python-pptx) — Modern karanlık / siber tema, başlık ve madde imleri.
3. Gelişmiş Kod & Doküman Üretimi
"""

import os
from pathlib import Path
from typing import List, Dict, Any, Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import pptx
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

DEFAULT_DESKTOP = Path(os.path.expanduser(r"~\Desktop"))


def create_excel_document(
    title: str,
    headers: List[str],
    rows: List[List[Any]],
    filename: Optional[str] = None,
    sheet_name: str = "Tablo",
    include_summary: bool = False
) -> Dict[str, Any]:
    """
    Kullanıcı komutlarıyla profesyonelce biçimlendirilmiş bir Excel çalışma kitabı oluşturur.
    """
    if not HAS_OPENPYXL:
        return {"status": "error", "message": "openpyxl kütüphanesi yüklü değil."}

    if not filename:
        clean_title = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).strip()
        filename = f"{clean_title or 'Jarvis_Tablo'}.xlsx"
    elif not filename.endswith(".xlsx"):
        filename += ".xlsx"

    out_path = DEFAULT_DESKTOP / filename

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Renk Paleti (Siber Teal / Kurumsal Lacivert)
    HEADER_FILL = PatternFill(start_color="0F3D3E", end_color="0F3D3E", fill_type="solid")
    HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    ROW_FILL_ALT = PatternFill(start_color="F2F7F7", end_color="F2F7F7", fill_type="solid")
    REG_FONT = Font(name="Segoe UI", size=10, color="1A1A1A")
    TITLE_FONT = Font(name="Segoe UI", size=15, bold=True, color="0F3D3E")

    thin_border_side = Side(border_style="thin", color="CCCCCC")
    BORDER = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
    ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
    ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

    # Başlık Satırı (A1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
    title_cell = ws.cell(row=1, column=1, value=title.upper())
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 34

    # Tablo Başlıkları (Row 3)
    start_row = 3
    ws.row_dimensions[start_row].height = 24
    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header_text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER

    # Veri Satırları
    current_row = start_row + 1
    for r_idx, row_data in enumerate(rows):
        ws.row_dimensions[current_row].height = 20
        use_alt = (r_idx % 2 == 1)
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = REG_FONT
            cell.border = BORDER
            if use_alt:
                cell.fill = ROW_FILL_ALT

            if isinstance(val, (int, float)):
                cell.alignment = ALIGN_RIGHT
                if isinstance(val, float):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = ALIGN_LEFT
        current_row += 1

    # Özet Satırı (Toplam)
    if include_summary and len(rows) > 0:
        ws.row_dimensions[current_row].height = 22
        sum_label = ws.cell(row=current_row, column=1, value="TOPLAM")
        sum_label.font = Font(name="Segoe UI", size=10, bold=True, color="0F3D3E")
        sum_label.border = BORDER
        for c in range(2, len(headers) + 1):
            col_letter = get_column_letter(c)
            # Sayısal değer kontrolü
            first_val = rows[0][c - 1] if len(rows[0]) >= c else None
            if isinstance(first_val, (int, float)):
                f_cell = ws.cell(
                    row=current_row, column=c,
                    value=f"=SUM({col_letter}{start_row + 1}:{col_letter}{current_row - 1})"
                )
                f_cell.font = Font(name="Segoe UI", size=10, bold=True)
                f_cell.border = BORDER
                f_cell.alignment = ALIGN_RIGHT
            else:
                empty = ws.cell(row=current_row, column=c, value="")
                empty.border = BORDER
        current_row += 1

    # Sütun Genişliklerini Otomatik Ayarla
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row == 1:
                continue  # Başlık satırını atla
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    try:
        wb.save(str(out_path))
        return {
            "status": "success",
            "message": f"Excel tablosu başarıyla oluşturuldu: {out_path.name}",
            "file_path": str(out_path),
            "rows_count": len(rows),
            "columns_count": len(headers)
        }
    except Exception as exc:
        return {"status": "error", "message": f"Excel dosyası kaydedilemedi: {exc}"}


def create_powerpoint_presentation(
    title: str,
    subtitle: str,
    slides_data: List[Dict[str, Any]],
    filename: Optional[str] = None
) -> Dict[str, Any]:
    """
    Modern koyu / siber temalı, tam teşekküllü PowerPoint slayt gösterisi oluşturur.
    slides_data formatı:
    [
        {"title": "Slayt Başlığı", "bullets": ["Madde 1", "Madde 2", "Madde 3"]},
        ...
    ]
    """
    if not HAS_PPTX:
        return {"status": "error", "message": "python-pptx kütüphanesi yüklü değil."}

    if not filename:
        clean_title = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).strip()
        filename = f"{clean_title or 'Jarvis_Sunum'}.pptx"
    elif not filename.endswith(".pptx"):
        filename += ".pptx"

    out_path = DEFAULT_DESKTOP / filename

    prs = Presentation()
    # 16:9 Geniş Ekran Formatı
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Renk Paleti (Jarvis Cyber Dark)
    COLOR_BG = RGBColor(6, 18, 20)          # Derin lacivert/siber siyah
    COLOR_CYAN = RGBColor(0, 240, 200)       # Vurgu Cam Göbeği
    COLOR_TEXT = RGBColor(235, 245, 245)     # Yumuşak Beyaz
    COLOR_MUTED = RGBColor(120, 160, 160)    # İkincil Metin
    COLOR_CARD = RGBColor(10, 28, 30)        # Kart arka planı

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(
            1,  # MSO_SHAPE.RECTANGLE
            0, 0, prs.slide_width, prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    # ── 1. BAŞLIK SLAYTI (TITLE SLIDE) ────────────────────────────────────────
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Başlık Metin Kutusu
    tb = slide1.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(11.0), Inches(2.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.name = "Segoe UI"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(20)
    p2.font.color.rgb = COLOR_MUTED
    p2.space_before = Pt(16)

    # Alt Köşe "Hazırlayan: J.A.R.V.I.S Autonomous AI"
    tb_footer = slide1.shapes.add_textbox(Inches(1.2), Inches(6.4), Inches(8.0), Inches(0.6))
    p_f = tb_footer.text_frame.paragraphs[0]
    p_f.text = "Hazırlayan: J.A.R.V.I.S Autonomous Intelligence Suite"
    p_f.font.name = "Segoe UI"
    p_f.font.size = Pt(11)
    p_f.font.color.rgb = COLOR_MUTED

    # ── 2. İÇERİK SLAYTLARI (CONTENT SLIDES) ──────────────────────────────────
    for idx, slide_info in enumerate(slides_data, start=1):
        slide = prs.slides.add_slide(blank_layout)
        set_slide_background(slide)

        # Kart Arka Planı (Hafif Açık Kutu)
        card = slide.shapes.add_shape(
            1,  # MSO_SHAPE.RECTANGLE
            Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.6)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = RGBColor(20, 50, 52)
        card.line.width = Pt(1.5)

        # Slayt Başlığı
        tb_title = slide.shapes.add_textbox(Inches(1.5), Inches(1.3), Inches(10.3), Inches(1.0))
        tf_title = tb_title.text_frame
        p_t = tf_title.paragraphs[0]
        s_title = slide_info.get("title", f"Bölüm {idx}")
        p_t.text = s_title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_CYAN

        # Madde İmleri (Bullets)
        tb_content = slide.shapes.add_textbox(Inches(1.5), Inches(2.4), Inches(10.3), Inches(3.8))
        tf_content = tb_content.text_frame
        tf_content.word_wrap = True

        bullets = slide_info.get("bullets", [])
        if isinstance(bullets, str):
            bullets = [bullets]

        for b_idx, bullet_text in enumerate(bullets):
            p_b = tf_content.paragraphs[0] if b_idx == 0 else tf_content.add_paragraph()
            p_b.text = f"•  {bullet_text}"
            p_b.font.name = "Segoe UI"
            p_b.font.size = Pt(16)
            p_b.font.color.rgb = COLOR_TEXT
            p_b.space_before = Pt(12)

        # Sağ Alt Slayt Numarası
        tb_num = slide.shapes.add_textbox(Inches(11.0), Inches(6.0), Inches(1.0), Inches(0.5))
        p_n = tb_num.text_frame.paragraphs[0]
        p_n.text = f"{idx + 1}"
        p_n.font.name = "Segoe UI"
        p_n.font.size = Pt(12)
        p_n.font.color.rgb = COLOR_MUTED
        p_n.alignment = PP_ALIGN.RIGHT

    try:
        prs.save(str(out_path))
        return {
            "status": "success",
            "message": f"PowerPoint sunumu başarıyla oluşturuldu: {out_path.name}",
            "file_path": str(out_path),
            "slides_count": len(slides_data) + 1
        }
    except Exception as exc:
        return {"status": "error", "message": f"Sunum kaydedilemedi: {exc}"}
